#!/usr/bin/env python3
"""Convert the China bonds vs equities report from Markdown to a styled Word file."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "reports" / "中国债券与股票资产比较研究报告.md"
DOCX = ROOT / "reports" / "中国债券与股票资产比较研究报告.docx"
FIG = ROOT / "reports" / "figures"

NAVY = RGBColor(0x1B, 0x36, 0x5D)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
BODY = RGBColor(0x1A, 0x1A, 0x1A)
SLATE = RGBColor(0x5A, 0x67, 0x7A)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=BODY, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_paragraph_spacing(p, before=0, after=8, line=1.15, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_hyperlink_style_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1B365D")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_mixed_runs(paragraph, text, size=11, color=BODY):
    """Parse **bold** and `code` in a line."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            set_run_font(run, size=size, color=color)
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=NAVY)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=size, color=TEAL)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def is_table_separator(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}", line))


def parse_table_row(line: str):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            text = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, before=2, after=2, line=1.08)
            run = p.add_run(re.sub(r"\*\*([^*]+)\*\*", r"\1", text))
            if i == 0:
                set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                set_cell_shading(cell, "1B365D")
            else:
                set_run_font(run, size=10, color=BODY)
                if i % 2 == 0:
                    set_cell_shading(cell, "F4F1EC")
                else:
                    set_cell_shading(cell, "FFFFFF")
    doc.add_paragraph()


def add_picture(doc, rel_path: str, caption: str | None = None):
    path = (MD.parent / rel_path).resolve()
    if not path.exists():
        path = FIG / Path(rel_path).name
    if not path.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[缺失配图：{rel_path}]")
        set_run_font(run, size=10, italic=True, color=SLATE)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=8, after=4)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.6))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(cap, before=0, after=10)
        r = cap.add_run(caption)
        set_run_font(r, size=9, italic=True, color=SLATE)


def convert():
    text = MD.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("中国债券与股票资产比较研究  ·  大类资产配置投研范式  ·  2026年8月")
    set_run_font(fr, size=8, color=SLATE)

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=4, after=8)
            add_hyperlink_style_border(p)
            i += 1
            continue

        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if img:
            add_picture(doc, img.group(2), caption=img.group(1) or None)
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=0 if first_h1 else 16, after=6)
            run = p.add_run(stripped[2:])
            set_run_font(run, size=22, bold=True, color=NAVY)
            first_h1 = False
            i += 1
            continue

        if stripped.startswith("**副标题"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=0, after=12)
            run = p.add_run(stripped.replace("**", ""))
            set_run_font(run, size=13, color=GOLD)
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=16, after=8)
            add_hyperlink_style_border(p)
            run = p.add_run(stripped[3:])
            set_run_font(run, size=16, bold=True, color=NAVY)
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=12, after=6)
            run = p.add_run(stripped[4:])
            set_run_font(run, size=13, bold=True, color=TEAL)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1].strip()):
            rows = [parse_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i].strip()))
                i += 1
            add_table(doc, rows)
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_spacing(p, before=1, after=3, line=1.15)
            add_mixed_runs(p, stripped[2:], size=11)
            i += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            set_paragraph_spacing(p, before=1, after=3, line=1.15)
            add_mixed_runs(p, numbered.group(2), size=11)
            i += 1
            continue

        # block formula
        if stripped.startswith("\\[") or stripped.endswith("\\]"):
            buf = [stripped]
            if not stripped.endswith("\\]"):
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("\\]"):
                    buf.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    buf.append(lines[i].strip())
            formula = " ".join(buf).replace("\\[", "").replace("\\]", "").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=6, after=6)
            run = p.add_run(formula)
            set_run_font(run, name="Cambria Math", size=12, italic=True, color=NAVY)
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=10, after=4)
            run = p.add_run(stripped.strip("*"))
            set_run_font(run, size=9, italic=True, color=SLATE)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        set_paragraph_spacing(p, before=2, after=6, line=1.22)
        add_mixed_runs(p, stripped, size=11)
        i += 1

    doc.save(DOCX)
    print(f"wrote {DOCX} ({DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    convert()
