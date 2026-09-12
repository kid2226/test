#!/usr/bin/env python3
"""Generate a Word copy of the financial-leasing industry review."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "中国金融租赁行业综述.md"
DOCX_PATH = ROOT / "中国金融租赁行业综述.docx"

CN_FONT = "WenQuanYi Micro Hei"
EN_FONT = "Calibri"
HEADING_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xC4, 0x5C, 0x26)


def set_run_font(run, *, size=11, bold=False, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = EN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_format(p, *, before=0, after=8, line=1.25, align=None, first_line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align is not None:
        p.alignment = align
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)


def add_runs_from_markdown(paragraph, text, *, size=11, color=None, base_bold=False):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=base_bold, color=color)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.25
    pf.space_after = Pt(8)

    for i, size, before, after in (
        (1, 18, 18, 10),
        (2, 15, 16, 8),
        (3, 13, 12, 6),
    ):
        st = styles[f"Heading {i}"]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = HEADING_COLOR
        st.font.name = EN_FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.15


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_format(p, before=2, after=2, line=1.15)
            text = row[j] if j < len(row) else ""
            add_runs_from_markdown(
                p,
                text,
                size=9,
                color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None,
                base_bold=(i == 0),
            )
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            if i == 0:
                shd.set(qn("w:fill"), "1F3A5F")
            elif i % 2 == 0:
                shd.set(qn("w:fill"), "F4F7FB")
            else:
                shd.set(qn("w:fill"), "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def parse_md_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows


def convert() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    configure_styles(doc)

    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            set_paragraph_format(p, before=0, after=2, line=1.1)
            run = p.add_run(line)
            set_run_font(run, size=10)
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            set_paragraph_format(p, before=4, after=4, line=1.0)
            run = p.add_run("─" * 28)
            set_run_font(run, size=10, color=RULE)
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_md_table(block))
            continue

        if re.match(r"^- ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_format(p, before=1, after=3, line=1.2)
            add_runs_from_markdown(p, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            set_paragraph_format(p, before=1, after=3, line=1.2)
            add_runs_from_markdown(p, re.sub(r"^\d+\.\s", "", stripped))
            i += 1
            continue

        if stripped.startswith("- ") and "：" in stripped[:20] or stripped.startswith("- 报告") or stripped.startswith("- 数据") or stripped.startswith("- 分析"):
            p = doc.add_paragraph()
            set_paragraph_format(p, before=0, after=4, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            add_runs_from_markdown(p, stripped[2:] if stripped.startswith("- ") else stripped, size=10.5, color=MUTED)
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            set_paragraph_format(p, before=8, after=4, line=1.2)
            add_runs_from_markdown(p, stripped, size=10, color=MUTED)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        # hanging meta lines like **副标题：**
        if stripped.startswith("**副标题"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(p, before=0, after=10, line=1.2)
            add_runs_from_markdown(p, stripped, size=12, color=HEADING_COLOR)
        else:
            set_paragraph_format(p, before=0, after=8, line=1.28, first_line=0.74)
            add_runs_from_markdown(p, stripped, size=11)
        i += 1

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("中国金融租赁行业综述  ·  2026年9月")
    set_run_font(run, size=8, color=MUTED)

    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    convert()
