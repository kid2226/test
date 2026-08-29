#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《我的生活疆域》主题演示文稿。"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from PIL import Image, ImageDraw, ImageFilter

ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets"
ASSET_DIR.mkdir(exist_ok=True)

# 16:9
W, H = Inches(13.333), Inches(7.5)

# 青绿宋韵 + 烟火暖色
INK = RGBColor(0x14, 0x28, 0x24)
PINE = RGBColor(0x1B, 0x4B, 0x40)
JADE = RGBColor(0x2A, 0x6E, 0x5C)
TEAL = RGBColor(0x3B, 0x8A, 0x74)
MINT = RGBColor(0xD8, 0xEB, 0xE3)
CREAM = RGBColor(0xF7, 0xF2, 0xE8)
IVORY = RGBColor(0xFF, 0xFB, 0xF5)
AMBER = RGBColor(0xC4, 0x8A, 0x38)
GOLD = RGBColor(0xE0, 0xC0, 0x7A)
CLAY = RGBColor(0xC2, 0x5B, 0x3C)
WARM = RGBColor(0xF3, 0xE4, 0xC8)
MUTED = RGBColor(0x5B, 0x6C, 0x66)
SOFT_INK = RGBColor(0x2C, 0x3E, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFC, 0xF7)


def _set_run_font(run, size, color, bold=False, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        for el in rPr.findall(qn(f"a:{tag}")):
            rPr.remove(el)
    latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", "Microsoft YaHei")
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)
    cs = etree.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", "Microsoft YaHei")


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    name="微软雅黑",
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, color, bold, name)
    return box


def add_para(tf, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, space_before=0, space_after=6):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, color, bold)
    return p


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def round_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    # 更利落的圆角
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def oval(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, page, total=12, light=False):
    color = RGBColor(0xA8, 0xB8, 0xB2) if light else RGBColor(0x8A, 0x9A, 0x94)
    add_text(
        slide,
        Inches(0.55),
        Inches(7.12),
        Inches(9.5),
        Inches(0.28),
        "我的生活疆域  ·  杭州市西湖区东山弄菜场",
        size=11,
        color=color,
    )
    add_text(
        slide,
        Inches(11.4),
        Inches(7.12),
        Inches(1.4),
        Inches(0.28),
        f"{page:02d}  /  {total:02d}",
        size=11,
        color=color,
        align=PP_ALIGN.RIGHT,
    )


def paint_mountains(path: Path, width=1600, height=520):
    """抽象青绿远山，封面用。"""
    img = Image.new("RGB", (width, height), (20, 40, 36))
    draw = ImageDraw.Draw(img, "RGBA")

    def mountain(points, color):
        draw.polygon(points, fill=color)

    mountain([(0, 360), (180, 210), (360, 340), (520, 160), (760, 330), (980, 140), (1200, 300), (1400, 180), (1600, 320), (1600, 520), (0, 520)], (27, 75, 64, 255))
    mountain([(0, 410), (220, 280), (430, 390), (640, 250), (880, 380), (1100, 230), (1320, 360), (1600, 270), (1600, 520), (0, 520)], (42, 110, 92, 230))
    mountain([(0, 460), (260, 350), (500, 440), (740, 320), (1020, 430), (1280, 340), (1600, 410), (1600, 520), (0, 520)], (59, 138, 116, 220))
    # 暖金薄雾
    for i in range(8):
        y = 300 + i * 18
        draw.rectangle((0, y, width, y + 10), fill=(224, 192, 122, 8 + i))
    img = img.filter(ImageFilter.SMOOTH_MORE)
    img.save(path)
    return path


def paint_icon(path: Path, kind: str, bg=(42, 110, 92), fg=(255, 251, 245)):
    s = 256
    img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((8, 8, s - 8, s - 8), fill=bg)
    c = fg
    if kind == "eye":
        d.ellipse((52, 92, 204, 168), outline=c, width=11)
        d.ellipse((102, 108, 154, 160), fill=c)
        d.ellipse((118, 122, 138, 142), fill=bg)
    elif kind == "ear":
        d.arc((86, 52, 186, 204), 240, 120, fill=c, width=12)
        d.arc((108, 84, 172, 176), 245, 115, fill=c, width=8)
        d.line([(88, 128), (112, 128)], fill=c, width=8)
    elif kind == "nose":
        d.arc((88, 58, 168, 188), 200, 340, fill=c, width=11)
        d.arc((96, 150, 128, 186), 20, 200, fill=c, width=8)
        d.arc((128, 150, 160, 186), 340, 160, fill=c, width=8)
        d.line([(108, 86), (148, 110)], fill=c, width=6)
        d.line([(148, 110), (148, 150)], fill=c, width=6)
    elif kind == "tongue":
        d.arc((70, 70, 186, 150), 200, 340, fill=c, width=11)
        d.chord((86, 118, 170, 202), 0, 180, outline=c, width=11)
        d.line([(128, 150), (128, 190)], fill=c, width=7)
    elif kind == "heart":
        d.ellipse((70, 78, 132, 140), fill=c)
        d.ellipse((124, 78, 186, 140), fill=c)
        d.polygon([(76, 118), (128, 190), (180, 118)], fill=c)
    elif kind == "map":
        d.polygon([(68, 72), (128, 56), (128, 196), (68, 180)], outline=c, width=9)
        d.polygon([(128, 56), (196, 78), (196, 198), (128, 196)], outline=c, width=9)
        d.line([(92, 108), (112, 100)], fill=c, width=6)
        d.line([(148, 96), (176, 108)], fill=c, width=6)
        d.ellipse((108, 118, 148, 158), outline=c, width=8)
        d.ellipse((122, 132, 134, 144), fill=c)
    elif kind == "ask":
        d.rounded_rectangle((72, 64, 184, 168), radius=28, outline=c, width=11)
        d.polygon([(104, 166), (88, 204), (136, 168)], fill=c)
        d.ellipse((110, 104, 128, 122), fill=c)
        d.ellipse((140, 104, 158, 122), fill=c)
    elif kind == "price":
        d.polygon([(86, 70), (186, 70), (186, 168), (136, 198), (86, 168)], outline=c, width=11)
        d.ellipse((118, 96, 154, 132), outline=c, width=8)
        d.line([(128, 132), (144, 160)], fill=c, width=8)
    elif kind == "fish":
        d.ellipse((48, 96, 176, 168), outline=c, width=10)
        d.polygon([(168, 118), (216, 86), (208, 130), (216, 174), (168, 146)], outline=c, width=9)
        d.ellipse((74, 118, 96, 140), fill=c)
        d.line([(120, 104), (132, 88)], fill=c, width=7)
    img.save(path)
    return path


def build_assets():
    mountains = paint_mountains(ASSET_DIR / "mountains.png")
    icons = {}
    for kind in ("eye", "ear", "nose", "tongue", "heart", "map", "ask", "fish", "price"):
        icons[kind] = paint_icon(ASSET_DIR / f"icon_{kind}.png", kind)
    icons["fish_warm"] = paint_icon(ASSET_DIR / "icon_fish_warm.png", "fish", bg=(194, 91, 60))
    icons["price_warm"] = paint_icon(ASSET_DIR / "icon_price_warm.png", "price", bg=(196, 138, 56))
    return mountains, icons


def content_header(slide, kicker, title, subtitle=None):
    rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.12), JADE)
    add_text(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.28), kicker, size=12, color=JADE, bold=True)
    add_text(slide, Inches(0.55), Inches(0.52), Inches(12.2), Inches(0.5), title, size=30, color=INK, bold=True)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(1.04), Inches(12.2), Inches(0.36), subtitle, size=14, color=MUTED)


def build():
    mountains, icons = build_assets()
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # ---------- 1 封面 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, INK)
    s.shapes.add_picture(str(mountains), Inches(0), Inches(4.55), Inches(13.333), Inches(2.95))
    rect(s, Inches(0), Inches(0), Inches(0.16), Inches(7.5), AMBER)
    add_text(s, Inches(0.7), Inches(0.55), Inches(8), Inches(0.32), "初一年级主题演讲", size=14, color=GOLD, bold=True)
    add_text(s, Inches(0.7), Inches(1.15), Inches(12), Inches(1.05), "我的生活疆域", size=52, color=WHITE, bold=True)
    add_text(s, Inches(0.7), Inches(2.25), Inches(11.5), Inches(0.5), "在东山弄，把生活画大一圈", size=24, color=GOLD)
    rect(s, Inches(0.7), Inches(2.92), Inches(1.35), Inches(0.05), AMBER)
    add_text(
        s,
        Inches(0.7),
        Inches(3.2),
        Inches(10.5),
        Inches(0.7),
        "一个初一新生的开疆记  ·  新空间 · 新尝试 · 新感受",
        size=16,
        color=MINT,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(6.85),
        Inches(10),
        Inches(0.35),
        "杭州市西湖区灵隐街道外东山弄86号   |   演讲人：________",
        size=13,
        color=GOLD,
    )
    notes(
        s,
        "开场：大家好，我是初一新生。今天分享的主题是“我的生活疆域”。我想讲一次不太一样的开疆之旅——去东山弄菜场。",
    )

    # ---------- 2 路线 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "OPENING  ·  01", "今天的开疆路线", "六步走完：从“什么是疆域”，走到“我把地图画大了一圈”。")
    steps = [
        ("01", "问疆域", "先想清楚\n生活版图是什么"),
        ("02", "选坐标", "为什么不去景区\n偏偏去菜场"),
        ("03", "敢尝试", "自己问路看价\n开口问摊主"),
        ("04", "用五感", "眼耳鼻舌心\n重新认识杭州"),
        ("05", "识人情", "烟火气里\n遇见一座城"),
        ("06", "画新图", "菜场在长大\n我也在长大"),
    ]
    left0 = 0.45
    gap = 2.12
    for i, (num, title, desc) in enumerate(steps):
        x = Inches(left0 + i * gap)
        round_rect(s, x, Inches(2.0), Inches(1.95), Inches(4.35), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        oval(s, x + Inches(0.62), Inches(2.28), Inches(0.7), Inches(0.7), JADE if i % 2 == 0 else CLAY)
        add_text(s, x + Inches(0.62), Inches(2.40), Inches(0.7), Inches(0.5), num, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, Inches(3.15), Inches(1.95), Inches(0.45), title, size=18, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.12), Inches(3.7), Inches(1.72), Inches(1.5), desc, size=13, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            rect(s, x + Inches(1.95), Inches(4.05), Inches(0.17), Inches(0.035), AMBER)
    footer(s, 2)
    notes(s, "先给同学看路线：今天不只是“我去逛了个菜场”，而是一次有地图、有尝试、有感受的开疆。")

    # ---------- 3 什么是疆域 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "THINK  ·  02", "疆域，不只是远方", "对初一的我来说，疆域不是国土，是我愿意伸手去碰的世界。")
    round_rect(s, Inches(0.5), Inches(1.7), Inches(6.0), Inches(4.95), PINE)
    add_text(s, Inches(0.85), Inches(2.0), Inches(5.3), Inches(0.35), "小学的生活地图", size=14, color=GOLD, bold=True)
    add_text(s, Inches(0.85), Inches(2.4), Inches(5.3), Inches(0.7), "又窄又安全", size=28, color=WHITE, bold=True)
    olds = [("家", "每天出发的原点"), ("学校", "最熟悉的半径"), ("小区小店", "偶尔的一点点远")]
    for i, (t, d) in enumerate(olds):
        y = Inches(3.3) + Inches(i * 0.95)
        oval(s, Inches(0.9), y, Inches(0.18), Inches(0.18), AMBER)
        add_text(s, Inches(1.25), y - Inches(0.08), Inches(4.8), Inches(0.35), t, size=18, color=WHITE, bold=True)
        add_text(s, Inches(1.25), y + Inches(0.28), Inches(4.8), Inches(0.32), d, size=13, color=MINT)
    cards = [
        ("敢去", "新空间：我愿意走进去的地方"),
        ("会看", "新观察：风景之外的烟火"),
        ("能感", "新感受：心里多出来的一圈"),
    ]
    for i, (t, d) in enumerate(cards):
        y = Inches(1.7) + Inches(i * 1.65)
        round_rect(s, Inches(6.8), y, Inches(6.0), Inches(1.5), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        oval(s, Inches(7.05), y + Inches(0.42), Inches(0.66), Inches(0.66), JADE)
        add_text(s, Inches(7.05), y + Inches(0.54), Inches(0.66), Inches(0.42), f"{i+1:02d}", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, Inches(7.95), y + Inches(0.32), Inches(4.5), Inches(0.45), t, size=22, color=INK, bold=True)
        add_text(s, Inches(7.95), y + Inches(0.82), Inches(4.5), Inches(0.4), d, size=14, color=MUTED)
    footer(s, 3)
    notes(s, "以前觉得疆域是历史书里的词。现在明白：它是我敢去、会看、能感受的地方。小学地图很窄，初一我想画大一圈。")

    # ---------- 4 坐标 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "PLACE  ·  03", "我的新坐标：东山弄", "西湖看风景，东山看烟火。离湖不到一千米，却是另一座杭州。")
    s.shapes.add_picture(str(icons["map"]), Inches(0.6), Inches(1.75), Inches(0.55), Inches(0.55))
    add_text(s, Inches(1.3), Inches(1.82), Inches(11), Inches(0.45), "杭州市西湖区灵隐街道外东山弄 86 号", size=22, color=INK, bold=True)
    facts = [
        ("不足千米", "距西湖直线距离", "步行大约 5 到 10 分钟"),
        ("二十多年", "附近居民的菜篮子", "辐射小区、高校数十万人"),
        ("东山集", "菜场的新名字", "从农贸市场长成邻里空间"),
        ("一菜一早", "这里的生活哲学", "买菜，也把早饭吃得像样"),
    ]
    for i, (t, a, b) in enumerate(facts):
        col, row = i % 2, i // 2
        x = Inches(0.55) + Inches(col * 6.35)
        y = Inches(2.6) + Inches(row * 2.05)
        round_rect(s, x, y, Inches(6.1), Inches(1.88), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        rect(s, x, y, Inches(0.12), Inches(1.88), JADE if i % 2 == 0 else CLAY)
        add_text(s, x + Inches(0.4), y + Inches(0.28), Inches(5.4), Inches(0.5), t, size=24, color=INK, bold=True)
        add_text(s, x + Inches(0.4), y + Inches(0.82), Inches(5.4), Inches(0.35), a, size=15, color=JADE, bold=True)
        add_text(s, x + Inches(0.4), y + Inches(1.2), Inches(5.4), Inches(0.35), b, size=14, color=MUTED)
    footer(s, 4)
    notes(s, "坐标：西湖区灵隐街道外东山弄86号。现在叫西湖·东山集。离西湖很近，却专门藏着烟火气。")

    # ---------- 5 为什么是菜场 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "WHY  ·  04", "别人去景区，我去菜场", "真正的城市生活，不在明信片上，而在买菜的人声里。")
    reasons = [
        ("反差", "景区是被安排好的美\n菜场是还在呼吸的生活", "我想看看风景背后的杭州"),
        ("成人礼", "第一次自己走进\n大人的世界", "问路、看价、开口，都是新尝试"),
        ("近处的远", "离家不远\n离旧自己却很远", "疆域不一定靠千里，靠勇气"),
    ]
    for i, (tag, title, desc) in enumerate(reasons):
        x = Inches(0.5) + Inches(i * 4.2)
        round_rect(s, x, Inches(1.75), Inches(3.95), Inches(4.85), PINE if i == 1 else CARD, line=None if i == 1 else RGBColor(0xE4, 0xD8, 0xC4))
        add_text(s, x + Inches(0.3), Inches(2.05), Inches(3.35), Inches(0.35), tag, size=13, color=GOLD if i == 1 else JADE, bold=True)
        add_text(s, x + Inches(0.3), Inches(2.55), Inches(3.35), Inches(1.6), title, size=22, color=WHITE if i == 1 else INK, bold=True)
        add_text(s, x + Inches(0.3), Inches(4.4), Inches(3.35), Inches(1.4), desc, size=15, color=MINT if i == 1 else MUTED)
    footer(s, 5)
    notes(s, "为什么不去西湖、不去博物馆？因为真正的城市生活藏在买菜的人声里。菜场是大人的世界，也是近处的远方。")

    # ---------- 6 三张地图 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "FRAME  ·  05", "一张地图，三种画法", "这是我给这次出行发明的方法：不记流水账，而画三张会呼吸的地图。")
    maps = [
        (icons["map"], JADE, "地理疆域", "空间被打开", "家—学校之外，多了一个坐标。一条弄堂，一座市集，离湖很近的人间。"),
        (icons["eye"], AMBER, "感官疆域", "感觉被打开", "用眼耳鼻舌心重新认识杭州：宋韵青绿，热油香，西湖鱼，生煎的脆底。"),
        (icons["heart"], CLAY, "人情疆域", "心里被打开", "摊主叫得出老客名字。我第一次被当成“会买菜的人”。"),
    ]
    for i, (icon, color, title, sub, desc) in enumerate(maps):
        x = Inches(0.5) + Inches(i * 4.2)
        round_rect(s, x, Inches(1.75), Inches(4.0), Inches(4.85), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        rect(s, x, Inches(1.75), Inches(4.0), Inches(0.12), color)
        s.shapes.add_picture(str(icon), x + Inches(0.3), Inches(2.1), Inches(0.62), Inches(0.62))
        add_text(s, x + Inches(0.3), Inches(2.9), Inches(3.4), Inches(0.45), title, size=22, color=INK, bold=True)
        add_text(s, x + Inches(0.3), Inches(3.4), Inches(3.4), Inches(0.4), sub, size=14, color=color, bold=True)
        add_text(s, x + Inches(0.3), Inches(4.0), Inches(3.4), Inches(2.0), desc, size=15, color=SOFT_INK)
    footer(s, 6)
    notes(s, "创新方法：不写流水账，画三张地图——地理、感官、人情。空间打开、感觉打开、心里打开。")

    # ---------- 7 新尝试 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "TRY  ·  06", "新空间里的三次出手", "紧张得手心出汗，也要把手伸出去。新尝试，本身就像开盲盒。")
    tries = [
        (icons["ask"], "自己问路", "第一次不用跟着大人走。弄堂、入口、摊区，都要自己认。"),
        (icons["price_warm"], "自己看价", "菜不是货架上的包装，是有季节、有产地、有故事的活物。"),
        (icons["fish_warm"], "自己开口", "问：“阿姨，这鱼是西湖里的吗？”摊主说：凌晨现捞，捞到什么卖什么。"),
    ]
    for i, (icon, title, desc) in enumerate(tries):
        y = Inches(1.7) + Inches(i * 1.6)
        round_rect(s, Inches(0.5), y, Inches(8.15), Inches(1.48), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        s.shapes.add_picture(str(icon), Inches(0.72), y + Inches(0.38), Inches(0.7), Inches(0.7))
        add_text(s, Inches(1.65), y + Inches(0.22), Inches(6.7), Inches(0.4), f"0{i+1}   {title}", size=20, color=INK, bold=True)
        add_text(s, Inches(1.65), y + Inches(0.7), Inches(6.7), Inches(0.58), desc, size=14, color=MUTED)
    round_rect(s, Inches(8.9), Inches(1.7), Inches(3.9), Inches(4.8), PINE)
    add_text(s, Inches(9.15), Inches(2.0), Inches(3.4), Inches(0.35), "关键隐喻", size=13, color=GOLD, bold=True)
    add_text(s, Inches(9.15), Inches(2.4), Inches(3.4), Inches(1.2), "西湖鱼\n像开盲盒", size=26, color=WHITE, bold=True)
    add_text(
        s,
        Inches(9.15),
        Inches(3.85),
        Inches(3.4),
        Inches(2.2),
        "全杭州几乎只有这里能买到正宗西湖鱼。凌晨现捞，白条、鲫鱼、步鱼，来什么卖什么。生活也是这样：你得先伸手，才会有惊喜。",
        size=14,
        color=MINT,
    )
    footer(s, 7)
    notes(s, "三个任务：问路、看价、开口。问西湖鱼时摊主说像开盲盒。新空间里的第一次尝试，本身就像开盲盒。")

    # ---------- 8 五感 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "SENSE  ·  07", "五感开疆记", "我把东山弄当成一座不用门票的博物馆，一件一件收藏。")
    senses = [
        (icons["eye"], "眼", "青绿撞上红鳞", "宋韵装修取自《千里江山图》，和活鱼青菜挤在一起，又雅又热闹。"),
        (icons["ear"], "耳", "杭州话在流动", "刀声、电子秤、有人喊：明虾很新鲜，给孙子带点回去！"),
        (icons["nose"], "鼻", "热油遇见咖啡", "葱姜、油墩儿，还有咖啡香。原来菜场也能边逛边喝一杯。"),
        (icons["tongue"], "舌", "一菜一早的幸福", "1.5元牛肉生煎底脆汁多，葱包桧儿、油墩儿，都是杭州的味道。"),
        (icons["heart"], "心", "烟火也可以温柔", "不是吵闹，是踏实。我第一次觉得：人间很近。"),
    ]
    for i, (icon, gan, title, desc) in enumerate(senses):
        x = Inches(0.38) + Inches(i * 2.58)
        round_rect(s, x, Inches(1.7), Inches(2.45), Inches(4.9), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        s.shapes.add_picture(str(icon), x + Inches(0.78), Inches(1.95), Inches(0.88), Inches(0.88))
        add_text(s, x, Inches(2.95), Inches(2.45), Inches(0.4), gan, size=20, color=JADE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.12), Inches(3.4), Inches(2.2), Inches(0.7), title, size=16, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.16), Inches(4.2), Inches(2.12), Inches(2.0), desc, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 8)
    notes(s, "五感：眼看青绿与活鱼，耳听杭州话，鼻闻热油和咖啡，舌尝生煎和葱包桧儿，心觉得烟火气也可以温柔。")

    # ---------- 9 人情 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "PEOPLE  ·  08", "人情，比菜更鲜", "地理地图好画，人情地图要用心记。")
    quotes = [
        ("摊主王大姐", "“老李，明虾很新鲜，要不要给孙子带点回去？”", "她叫得出老顾客的名字，连谁家添了孩子都知道。"),
        ("买鱼的华阿姨", "“西湖步鱼配春笋雪菜炖汤，鲜掉眉毛。”", "我站在旁边听，像偷师一门只有杭州人才懂的功课。"),
        ("初一的我", "“阿姨，这鱼是西湖里的吗？”", "第一次被当成会买菜的人，而不是跟着走的小孩。"),
    ]
    for i, (who, quote, feel) in enumerate(quotes):
        y = Inches(1.68) + Inches(i * 1.68)
        round_rect(s, Inches(0.5), y, Inches(12.3), Inches(1.54), PINE if i == 2 else CARD, line=None if i == 2 else RGBColor(0xE4, 0xD8, 0xC4))
        add_text(s, Inches(0.85), y + Inches(0.18), Inches(11.6), Inches(0.32), who, size=13, color=GOLD if i == 2 else JADE, bold=True)
        add_text(s, Inches(0.85), y + Inches(0.5), Inches(11.6), Inches(0.42), quote, size=18, color=WHITE if i == 2 else INK, bold=True)
        add_text(s, Inches(0.85), y + Inches(1.0), Inches(11.6), Inches(0.35), feel, size=14, color=MINT if i == 2 else MUTED)
    footer(s, 9)
    notes(s, "人情比菜更鲜。摊主认识老客，华阿姨教烧鱼。我第一次被当成会买菜的人。疆域是用眼睛看、用嘴巴问、用心记住的。")

    # ---------- 10 镜像 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "MIRROR  ·  09", "菜场在长大，我也在长大", "最巧的发现：东山弄自己也在拓宽疆域。我和它，好像在做同一件事。")
    # 时间轴
    rect(s, Inches(1.1), Inches(2.55), Inches(11.1), Inches(0.045), AMBER)
    nodes = [
        ("过去", "传统农贸市场", "买菜、称重、回家"),
        ("现在", "一菜一早", "小吃、早餐、西湖鱼"),
        ("更新", "邻里共生", "咖啡、便民、诗歌展"),
        ("对照", "初一的我", "敢问、敢试、敢观察"),
    ]
    for i, (a, b, c) in enumerate(nodes):
        x = Inches(0.7) + Inches(i * 3.15)
        oval(s, x + Inches(1.05), Inches(2.38), Inches(0.38), Inches(0.38), CLAY if i == 3 else JADE)
        add_text(s, x, Inches(2.9), Inches(2.5), Inches(0.35), a, size=13, color=JADE if i < 3 else CLAY, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, Inches(3.25), Inches(2.5), Inches(0.45), b, size=18, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, Inches(3.75), Inches(2.5), Inches(0.7), c, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    round_rect(s, Inches(0.5), Inches(4.7), Inches(12.3), Inches(1.9), PINE)
    add_text(
        s,
        Inches(0.9),
        Inches(5.0),
        Inches(11.5),
        Inches(1.35),
        "黑珍珠餐厅的大厨来开小吃店，好味道变得人人吃得起；诗歌展办到菜场里，烟火气撞上文学。一座老菜场把世界变大，一个初一新生也把世界变大。我们都让自己的疆域，更大一点，也更有温度一点。",
        size=16,
        color=WHITE,
    )
    footer(s, 10)
    notes(s, "镜像：菜场从农贸市场变成一菜一早、邻里共生，甚至有大厨和诗歌展。菜场在扩疆，我也在扩疆。")

    # ---------- 11 新版图 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    content_header(s, "GROW  ·  10", "走出来，版图大了一圈", "地理上多了一条弄堂，能力上多了勇气，理解上多了一座人间。")
    gains = [
        ("01 地理", "多一个坐标", "家和学校之外，东山弄成了我认识杭州的新入口。"),
        ("02 能力", "敢问敢试", "问路、看价、开口。紧张可以有，退缩不可以。"),
        ("03 理解", "城是人间", "城市不只是风景，更是摊主、阿姨、热气和招呼。"),
        ("04 下一站", "继续画大", "下一次，我还想走进更新的空间，做更新的尝试。"),
    ]
    for i, (k, t, d) in enumerate(gains):
        col, row = i % 2, i // 2
        x = Inches(0.5) + Inches(col * 6.4)
        y = Inches(1.7) + Inches(row * 2.45)
        round_rect(s, x, y, Inches(6.15), Inches(2.25), CARD, line=RGBColor(0xE4, 0xD8, 0xC4))
        add_text(s, x + Inches(0.35), y + Inches(0.28), Inches(5.5), Inches(0.35), k, size=13, color=JADE, bold=True)
        add_text(s, x + Inches(0.35), y + Inches(0.68), Inches(5.5), Inches(0.45), t, size=24, color=INK, bold=True)
        add_text(s, x + Inches(0.35), y + Inches(1.25), Inches(5.5), Inches(0.7), d, size=15, color=MUTED)
    footer(s, 11)
    notes(s, "收获：地理多一个坐标，能力上敢问敢试，理解上知道城市是人间。下一站还要把地图画大。")

    # ---------- 12 结束 ----------
    s = prs.slides.add_slide(blank)
    set_bg(s, INK)
    s.shapes.add_picture(str(mountains), Inches(0), Inches(4.7), Inches(13.333), Inches(2.8))
    rect(s, Inches(0), Inches(0), Inches(0.16), Inches(7.5), AMBER)
    add_text(s, Inches(0.7), Inches(0.7), Inches(11), Inches(0.35), "CLOSING", size=13, color=GOLD, bold=True)
    add_text(s, Inches(0.7), Inches(1.25), Inches(12), Inches(1.8), "疆域不在远方，\n在我迈出的第一步。", size=36, color=WHITE, bold=True)
    rect(s, Inches(0.7), Inches(3.3), Inches(1.35), Inches(0.05), AMBER)
    add_text(
        s,
        Inches(0.7),
        Inches(3.55),
        Inches(11.5),
        Inches(0.8),
        "走进新的空间，做新的尝试，就会有特别的感受。\n谢谢老师，谢谢同学。",
        size=18,
        color=MINT,
    )
    add_text(s, Inches(0.7), Inches(6.9), Inches(10), Inches(0.3), "我的生活疆域  ·  东山弄", size=13, color=GOLD)
    notes(s, "收束金句：疆域不在远方，在我迈出的第一步。谢谢老师，谢谢同学。")

    out = ROOT / "我的生活疆域.pptx"
    prs.save(out)
    print(f"saved: {out}")
    return out


if __name__ == "__main__":
    build()
