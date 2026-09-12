#!/usr/bin/env python3
"""Convert the EV battery markdown briefing into a styled Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "研究报告_电动汽车电池市场投资分析.md"
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUT_DIR / "电动汽车电池市场投资分析.docx"

NAVY = "17365D"
BLUE = "2F75B5"
DARK = "1D2939"
MID_GRAY = "667085"
LIGHT_BLUE = "DCEAF7"
ROW_ALT = "F8FAFC"
QUOTE_BG = "F4F6F8"
BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
LATIN_FONT = "Calibri"

INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def set_run_font(run, name=BODY_FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), LATIN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    rpr = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    rpr.append(begin)
    rpr.append(instr)
    rpr.append(end)
    set_run_font(run, HEAD_FONT, size=9, color=MID_GRAY)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)


def configure_page(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("电动汽车电池市场投资分析  ·  仅供研究讨论")
    set_run_font(run, HEAD_FONT, size=8.5, color=MID_GRAY)
    hp.paragraph_format.space_after = Pt(2)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = fp.add_run("不构成投资建议    ·    ")
    set_run_font(left, HEAD_FONT, size=8.5, color=MID_GRAY)
    add_page_field(fp)


def add_bottom_border(paragraph, color=NAVY, size="12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_left_border(paragraph, color=BLUE, size="18") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), color)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def split_table_cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def parse_inline(paragraph, text: str, size=11, color=DARK, name=BODY_FONT) -> None:
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, name, size=size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name, size=size, italic=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", size=size - 0.5, color=NAVY)
            run.font.name = "Consolas"
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, _url = part[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            set_run_font(run, name, size=size, color=BLUE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, name, size=size, color=color)


def add_paragraph(doc, text: str, *, first_line=True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.35
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    parse_inline(p, text, size=11)
    return p


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph()
    keep_with_next(p)
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(text)
        set_run_font(run, HEAD_FONT, size=16, bold=True, color=NAVY)
        add_bottom_border(p, NAVY, "12")
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, HEAD_FONT, size=13, bold=True, color=BLUE)
    else:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_run_font(run, HEAD_FONT, size=12, bold=True, color=NAVY)


def add_list_item(doc, text: str, ordered: bool) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    parse_inline(p, text, size=11)


def add_quote(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.3
    add_left_border(p, BLUE, "18")
    parse_inline(p, text, size=10.5, color=MID_GRAY)


def add_caption(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, HEAD_FONT, size=9.5, bold=True, color=MID_GRAY)


def add_image(doc, rel_path: str) -> None:
    path = (ROOT / rel_path).resolve()
    if not path.exists():
        add_paragraph(doc, f"[缺失图片：{rel_path}]", first_line=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(16.2))


def add_code_block(doc, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.3)
    run = p.add_run("\n".join(lines))
    set_run_font(run, "Consolas", size=9.5, color=NAVY)
    run.font.name = "Consolas"


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    usable = 16.6
    # Give the last column more room when there are many columns
    if ncols >= 6:
        widths = [1.4, 3.4, 1.6, 2.4, 2.8, 3.0][:ncols]
        while len(widths) < ncols:
            widths.append(usable / ncols)
        scale = usable / sum(widths)
        widths = [w * scale for w in widths]
    elif ncols == 4:
        widths = [2.0, 4.2, 4.2, 6.2]
    elif ncols == 3:
        widths = [3.6, 6.5, 6.5]
    else:
        widths = [usable / ncols] * ncols

    for i, row in enumerate(rows):
        tr = table.rows[i]
        prevent_row_split(tr)
        if i == 0:
            set_repeat_header(tr)
        for j in range(ncols):
            cell = tr.cells[j]
            cell.width = Cm(widths[j])
            set_cell_margins(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            text = row[j] if j < len(row) else ""
            parse_inline(
                p,
                text,
                size=8.5 if ncols >= 6 else 9,
                color="FFFFFF" if i == 0 else DARK,
                name=HEAD_FONT if i == 0 else BODY_FONT,
            )
            if i == 0:
                shade_cell(cell, NAVY)
                for run in p.runs:
                    run.bold = True
            elif i % 2 == 0:
                shade_cell(cell, ROW_ALT)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def add_title_block(doc, title: str, metas: list[str]) -> None:
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
    eyebrow.paragraph_format.space_after = Pt(6)
    r = eyebrow.add_run("研究报告")
    set_run_font(r, HEAD_FONT, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(title)
    set_run_font(run, HEAD_FONT, size=22, bold=True, color=NAVY)
    add_bottom_border(p, NAVY, "18")

    for meta in metas:
        mp = doc.add_paragraph()
        mp.paragraph_format.space_after = Pt(2)
        mp.paragraph_format.space_before = Pt(0)
        parse_inline(mp, meta, size=10.5, color=MID_GRAY, name=HEAD_FONT)

    blank = doc.add_paragraph()
    blank.paragraph_format.space_after = Pt(8)


def convert() -> Path:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_page(section)
    add_header_footer(section)

    i = 0
    metas: list[str] = []

    # Title + metadata until first ---
    if lines and lines[0].startswith("# "):
        title = lines[0][2:].strip()
        i = 1
        while i < len(lines) and lines[i].strip() != "---":
            raw = lines[i].strip()
            if raw:
                metas.append(raw.rstrip("  "))
            i += 1
        if i < len(lines) and lines[i].strip() == "---":
            i += 1
        add_title_block(doc, title, metas)

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("```"):
            fence = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                fence.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, fence)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            rows = [split_table_cells(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_cells(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if img:
            add_image(doc, img.group(2))
            i += 1
            continue

        if stripped.startswith("**图") and stripped.endswith("**"):
            add_caption(doc, stripped.strip("*").replace("　", " "))
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 3)
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, stripped[3:], 2)
            i += 1
            continue
        if stripped.startswith("# "):
            add_heading(doc, stripped[2:], 1)
            i += 1
            continue

        if stripped.startswith("> "):
            add_quote(doc, stripped[2:])
            i += 1
            continue

        m_ul = re.match(r"^[-*] (.+)$", stripped)
        if m_ul:
            add_list_item(doc, m_ul.group(1), False)
            i += 1
            continue

        m_ol = re.match(r"^(\d+)\. (.+)$", stripped)
        if m_ol:
            add_list_item(doc, f"{m_ol.group(1)}. {m_ol.group(2)}", True)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        add_paragraph(doc, stripped, first_line=True)
        i += 1

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = convert()
    print(f"wrote {path} ({path.stat().st_size} bytes)")
